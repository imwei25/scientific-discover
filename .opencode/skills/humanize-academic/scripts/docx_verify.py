#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""就地改写的校验闸：证明"只改了字，没碰别的"，并核对数字与引用没被动过。

四道闸：
  A. zip 层    条目集合一致；没被改的部件【逐字节相同】（图片、样式、主题、字体表…）
  B. 结构层    表/行/单元格/合并格/图/域/公式/分节/段落数一一相等
  C. 文本层    润色档：修订模式下"拒绝全部修订"应当**还原成原文**（没有静默丢字的机械证明）
               翻译档：**逐段点名漏译**——译文与原文一字不差、且原文含源语言字符的段落
  D. 不变量    数字、引用标记、DOI/PMID、（可选）术语在前后一致（复用 check_invariants 的正则）

润色档与翻译档必须分开，因为两者的"正常"完全不同：翻译时每段都变、篇幅涨 1.5–2 倍，
拿润色那套闸去卡会满屏假阳性；反过来，润色不需要查漏译，翻译却最怕模型偷偷跳段。

用法：
  python docx_verify.py 原稿.docx 改后.docx
  python docx_verify.py 原稿.docx 改后.docx --terms "HFpEF,SGLT2i,eGFR"
  python docx_verify.py 原稿.docx 译后.docx --mode translate --expect-lang en

退出码：有硬伤=1，全过=0。
"""
import argparse
import re
import sys
import zipfile
from collections import Counter

import check_invariants as CI
import docx_ooxml as X

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# 结构指纹：这些标签的数量一旦变了，就说明改写动到了文档结构而不只是文字。
STRUCT_TAGS = [("tbl", "表"), ("tr", "表格行"), ("tc", "单元格"),
               ("vMerge", "纵向合并格"), ("gridSpan", "横向合并格"),
               ("drawing", "图"), ("instrText", "域代码"), ("fldChar", "域标记"),
               ("sectPr", "分节"), ("p", "段落"), ("footnoteReference", "脚注引用"),
               ("hyperlink", "超链接"), ("bookmarkStart", "书签")]


def text_view(root, mode):
    """accepted=接受全部修订后的文字；rejected=拒绝全部修订（即还原成原文）。

    段与段之间插一个换行——直接拼接会让上一段的结尾和下一段的开头黏成一个 token。
    实测踩过：某段以 DOI `…-00789` 收尾、下一段以中文起头，译成英文后拼出
    `00789English`，数字正则因为后面紧跟字母而整个匹配失败，报"丢失 00789"。
    文档里那个数字一直好好的，是校验器自己造的假警报。
    """
    out = []
    for el in root.iter():
        if el.tag == X.W + "p":
            out.append("\n")
        elif el.tag == X.W + "t":
            in_ins = any(a.tag == X.W + "ins" for a in el.iterancestors())
            if mode == "accepted" or not in_ins:
                out.append(el.text or "")
        elif el.tag == X.W + "delText" and mode == "rejected":
            out.append(el.text or "")
    return "".join(out)


def all_text(path, mode):
    return "".join(text_view(X.load(path, part), mode) for part in X.parts_of(path))


CJK_RE = re.compile(r"[一-鿿]")
LATIN_WORD_RE = re.compile(r"[A-Za-z]{3,}")


def para_pairs(before, after):
    """按部件与段落序号一一配对（结构守恒，所以序号天然对齐）。"""
    for part in X.parts_of(before):
        ra, rb = X.load(before, part), X.load(after, part)
        pa, pb = X.paragraphs(ra), X.paragraphs(rb)
        if len(pa) != len(pb):
            continue                      # 段落数不等已由 B 闸报出，这里不重复报
        prefix = X.part_prefix(part)
        for i, (x, y) in enumerate(zip(pa, pb)):
            yield f"{prefix}{i:04d}", x, y


def editable_text(p):
    return "".join(X.run_text(r) for k, v in X.segments(p) if k == "edit" for r in v)


def check_translation(before, after, expect):
    """翻译档的 C 闸：漏译点名 + 残留源语言统计。

    漏译是全文翻译最常见的事故（模型跳段、或把"这段不用翻"自作主张地执行），
    而它在成品里几乎看不出来——一份 40 页的译稿里夹着 3 段中文，用户翻到才发现。
    判据取"译文与原文一字不差 **且** 原文里有源语言字符"，不拿"含中文"单独判：
    英文稿里本来就可能引用中文机构名，那不是漏译。

    两处都只看【可改文字】，把 ⟦冻结片⟧（域结果、公式、图内文字）排除在外——
    那些文字是**故意**保持原样的，算进来的话，每一份带中文引文域的稿子都会被
    误判成"半译"，警报一多用户就不看了。
    """
    residual = CJK_RE if expect.startswith("en") else LATIN_WORD_RE
    missed, resid = [], []
    for pid, pa, pb in para_pairs(before, after):
        old, new = editable_text(pa), editable_text(pb)
        if not old.strip():
            continue                      # 整段都是域/公式：本来就没有可翻的字
        if old == new and residual.search(old):
            missed.append((pid, old))
        elif residual.search(new) and len(residual.findall(new)) >= 4:
            resid.append((pid, new))
    return missed, resid


def struct(path):
    c = Counter()
    for part in X.parts_of(path):
        root = X.load(path, part)
        for tag, _ in STRUCT_TAGS:
            c[tag] += len(root.findall(".//" + X.W + tag))
    return c


def main():
    ap = argparse.ArgumentParser(description="docx 就地改写校验")
    ap.add_argument("before"); ap.add_argument("after")
    ap.add_argument("--mode", choices=["polish", "translate"], default="polish",
                    help="polish=润色档（查修订完整性与篇幅）；translate=翻译档（查漏译与残留）")
    ap.add_argument("--expect-lang", default="en",
                    help="翻译档的目标语言（en / zh），决定拿什么字符判残留")
    ap.add_argument("--terms", default="", help="逗号分隔的术语白名单，逐个计数比对"
                    "（翻译档里用来钉住不该被翻译的基因名/缩写）")
    ap.add_argument("--auto-terms", action="store_true",
                    help="自动把基因名/化学式(TP53/SGLT2/IL-6)纳入比对")
    args = ap.parse_args()
    fail, warn = [], []

    za, zb = zipfile.ZipFile(args.before), zipfile.ZipFile(args.after)
    na, nb = set(za.namelist()), set(zb.namelist())
    print("== A. zip 层 ==")
    if na != nb:
        fail.append(f"条目集合不同：多了 {sorted(nb - na)[:3]}，少了 {sorted(na - nb)[:3]}")
    same, diff = [], []
    for n in sorted(na & nb):
        (same if za.read(n) == zb.read(n) else diff).append(n)
    print(f"  条目 {len(na)} 个：逐字节相同 {len(same)} 个，被改动 {len(diff)} 个"
          f"（{', '.join(d.replace('word/', '') for d in diff) or '无'}）")
    for n in diff:
        if not X.PART_RE.match(n):
            fail.append(f"不该被改的条目被改了：{n}")

    print("== B. 结构层 ==")
    sa, sb = struct(args.before), struct(args.after)
    for tag, name in STRUCT_TAGS:
        mark = "" if sa[tag] == sb[tag] else "  ✗"
        if sa[tag] or sb[tag]:
            print(f"  {name}：{sa[tag]} → {sb[tag]}{mark}")
        if sa[tag] != sb[tag]:
            fail.append(f"{name}数量变了 {sa[tag]}→{sb[tag]}")
    try:
        import docx
        docx.Document(args.after)
        print("  python-docx 打开产物：成功")
    except Exception as e:
        fail.append(f"python-docx 打不开产物：{e}")

    print("== C. 文本层 ==")
    orig = all_text(args.before, "accepted")
    rejected = all_text(args.after, "rejected")
    accepted = all_text(args.after, "accepted")
    tracked = rejected != accepted
    print(f"  原文 {len(orig)} 字 → 改后 {len(accepted)} 字"
          f"（{(len(accepted) - len(orig)) / max(len(orig), 1) * 100:+.1f}%）")

    if args.mode == "polish":
        print(f"  产物是否带修订标记：{'是' if tracked else '否（直接覆盖）'}")
        if tracked:
            if rejected == orig:
                print("  拒绝全部修订 → 与原文逐字相同 ✔（没有任何改动绕过修订标记）")
            else:
                fail.append("拒绝全部修订后与原文不一致——有文字被改却没有留下修订标记，"
                            "用户在 Word 里看不见它，也没法拒绝")
        # 篇幅这一条只对润色成立：翻译时中译英涨 1.5–2 倍是正常的
        if abs(len(accepted) - len(orig)) > 0.15 * max(len(orig), 1):
            warn.append("篇幅变动超过 ±15%，润色不该改这么多，请核对是否有整段被删或被扩写")
    else:
        missed, resid = check_translation(args.before, args.after, args.expect_lang)
        if missed:
            fail.append(f"有 {len(missed)} 段漏译（与原文一字不差）")
            print(f"  [!] 漏译 {len(missed)} 段——与原文一字不差：")
            for pid, txt in missed[:10]:
                print(f"      {pid}：{txt[:34]}…")
            if len(missed) > 10:
                print(f"      …另有 {len(missed) - 10} 段")
        else:
            print("  漏译检查：没有任何段落与原文一字不差 ✔")
        if resid:
            warn.append(f"{len(resid)} 段译文里仍有较多源语言字符，请人工看一眼是否半译")
            print(f"  [~] {len(resid)} 段译文里仍有成片的源语言字符（可能是半译，也可能是"
                  f"机构名/专有名词的正当保留）：")
            for pid, txt in resid[:5]:
                print(f"      {pid}：{txt[:34]}…")
        # 版面提醒：就地改写保留原列宽，长出来的译文只能靠换行消化
        ratio = len(accepted) / max(len(orig), 1)
        if ratio > 1.3:
            warn.append(f"译文比原文长 {(ratio - 1) * 100:.0f}%——原表格列宽是固定的，"
                        f"窄列里的长句会把行撑高，页数也会变，交付时要说一声")

    print("== D. 不变量 ==")
    eb, ea = {}, {}
    eb["数字"] = Counter(m.group(0).strip() for m in CI.NUM_RE.finditer(orig))
    ea["数字"] = Counter(m.group(0).strip() for m in CI.NUM_RE.finditer(accepted))
    # DOI 用【包含】判而不是 token 相等：参考文献表里 DOI 后面常常紧跟正文没有空格，
    # 正则会把后面的字母一起吃进来（`10.3389/fendo.2018.00432.本文` → 中文停住，
    # 但译成英文后变成 `...00432.This`）。token 比对在这种稿子上必然假红，
    # 而 DOI 真正要保证的只有一件事：原文里的每个 DOI 在产物里还在。
    lost_doi = sorted({m.group(0).rstrip(".,;)") for m in CI.DOI_RE.finditer(orig)}
                      - {d for d in [m.group(0).rstrip(".,;)") for m in CI.DOI_RE.finditer(accepted)]})
    lost_doi = [d for d in lost_doi if d not in accepted]
    for name, rx in (("引用[n]", CI.BRACKET_CITE_RE), ("PMID", CI.PMID_RE)):
        eb[name] = Counter(rx.findall(orig)); ea[name] = Counter(rx.findall(accepted))
    # 作者-年引用在【翻译】里只能比条数，不能比字符串：「（中泰证券，2025）」译成
    # 「(Zhongtai Securities, 2025)」是正确的翻译，不是丢引用。而条数少了就是真丢了。
    # 润色档仍按字符串严比——那里作者名本来就不该变。
    cite_b = CI.PAREN_CITE_RE.findall(orig) + CI.CJK_CITE_RE.findall(orig)
    cite_a = CI.PAREN_CITE_RE.findall(accepted) + CI.CJK_CITE_RE.findall(accepted)
    if args.mode == "translate":
        if len(cite_a) < len(cite_b):
            fail.append(f"作者-年引用少了 {len(cite_b) - len(cite_a)} 条"
                        f"（{len(cite_b)} → {len(cite_a)}）——译文可以换写法，但不许整条丢")
        elif cite_b:
            print(f"  作者-年引用 {len(cite_b)} → {len(cite_a)} 条（只比条数：作者名随语种变属正常）")
    else:
        eb["引用(作者,年)"] = Counter(cite_b)
        ea["引用(作者,年)"] = Counter(cite_a)
    terms = {t.strip() for t in args.terms.split(",") if t.strip()}
    if args.auto_terms:
        terms |= set(CI.AUTOTERM_RE.findall(orig)) | set(CI.AUTOTERM_RE.findall(accepted))
    if terms:
        eb["术语"] = Counter({t: orig.count(t) for t in terms})
        ea["术语"] = Counter({t: accepted.count(t) for t in terms})

    dirty = 0
    if lost_doi:
        dirty += 1
        fail.append(f"{len(lost_doi)} 个 DOI 在产物里找不到了")
        print("  [!] DOI 丢失：")
        for d in lost_doi[:8]:
            print(f"      - {d}")
    for kind in eb:
        lost, added = eb[kind] - ea[kind], ea[kind] - eb[kind]
        if not lost and not added:
            continue
        dirty += 1
        print(f"  [!] {kind} 有变化：")
        for k, n in list(lost.items())[:8]:
            print(f"      - 丢失 {k!r} ×{n}")
        for k, n in list(added.items())[:8]:
            print(f"      + 新增 {k!r} ×{n}")
        if kind in ("数字", "DOI", "PMID", "引用[n]", "引用(作者,年)"):
            fail.append(f"{kind}在改写前后不一致——润色不许动数据与引用")
    if not dirty:
        print("  数字、引用标记、DOI/PMID" + ("、术语" if terms else "") + " 前后一致 ✔"
              + (f"（DOI {len(set(m.group(0) for m in CI.DOI_RE.finditer(orig)))} 个全在）"
                 if CI.DOI_RE.search(orig) else ""))

    print("-" * 56)
    for w in warn:
        print(f"[WARN] {w}")
    if fail:
        for f in fail:
            print(f"[FAIL] {f}")
        print("\n这些是硬伤：修好再交付，别拿这份文件出件。")
        sys.exit(1)
    print("[OK] 结构、格式载体、数字与引用均未被改动；改的只有文字本身。")


if __name__ == "__main__":
    main()
