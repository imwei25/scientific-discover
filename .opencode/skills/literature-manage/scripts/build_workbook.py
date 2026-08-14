#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 library.json（模型填好的文献台账）出成【多 sheet】的 library.xlsx —— 一类一个 sheet。

library.json 的形状（本技能全流程的真值源，界面上的"改分类""按分类归档"也都读它）：
{
  "classified": true,                       // false = 不分类，全部放一个 sheet
  "rule": "按研究类型分",                   // 分类标准原话（用户给的；没给就写 AI 自定的口径）
  "columns": ["文件名","标题","年份","作者","杂志","核心观点"],
  "fields":  ["file","title","year","authors","journal","point"],
  "categories": ["随机对照试验","队列研究"],  // sheet 顺序；缺省按 records 里首次出现的顺序
  "records": [
    {"file":"a.pdf","category":"随机对照试验","title":"...","year":"2021",
     "authors":"Smith J, et al.","journal":"Lancet","point":"一句话说清它的核心结论"}
  ]
}

columns/fields 是【每次自己定的】，不是写死的六列：整理的不是文献（笔记、报告、会议记录、
草稿、说明文档……）时，作者 / 杂志 / 核心观点没有对应物，应改成
  "columns": ["文件名","标题","日期","主要内容"], "fields": ["file","title","date","summary"]
文献与非文献混在一个文件夹里，就在文献六列后面补一列「主要内容」，各行填自己有的那一列
（另一边留空或写「不适用」）。

铁律：抽不到的字段一律写「原文未标注」，**不许猜、不许用背景知识补**（年份、杂志尤其容易被脑补）。

用法：
  python build_workbook.py                       # 读当前目录 library.json → library.xlsx
  python build_workbook.py --json library.json --out library.xlsx
"""
import argparse
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少 openpyxl：先跑 env-setup 技能")

DEFAULT_COLUMNS = ["文件名", "标题", "年份", "作者", "杂志", "核心观点"]
DEFAULT_FIELDS = ["file", "title", "year", "authors", "journal", "point"]
UNCLASSIFIED = "未分类"
# 每列宽度（字符）。核心观点 / 主要内容最宽并自动换行——这一列才是用户真正要读的东西。
# 「主要内容」是【非文献】文件（笔记、报告、记录、说明……）那一列：它们没有作者与杂志可写，
# 强行留着那两列只会得到一整列「原文未标注」。
WIDTHS = {"文件名": 34, "标题": 42, "年份": 8, "作者": 22, "杂志": 22, "核心观点": 72,
          "主要内容": 72, "类型": 12, "日期": 12, "备注": 30}


def sheet_name(raw, used):
    """Excel 的 sheet 名限制：≤31 字符、不能含 []:*?/\\、不能重名、不能为空"""
    name = re.sub(r"[\[\]:*?/\\]", "_", str(raw or "").strip()) or UNCLASSIFIED
    name = name[:31]
    if name.lower() in used:
        for i in range(2, 100):
            cand = f"{name[:28]}({i})"
            if cand.lower() not in used:
                name = cand
                break
    used.add(name.lower())
    return name


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", default="library.json")
    ap.add_argument("--out", default="library.xlsx")
    a = ap.parse_args()

    if not os.path.exists(a.json):
        sys.exit(f"找不到 {a.json} —— 先读 library_texts/ 里的正文，把台账写成 library.json 再来")
    with open(a.json, encoding="utf-8") as f:
        lib = json.load(f)

    columns = lib.get("columns") or DEFAULT_COLUMNS
    fields = lib.get("fields") or DEFAULT_FIELDS
    if len(columns) != len(fields):
        sys.exit("library.json 里 columns 与 fields 长度对不上（一个表头对一个字段名）")
    records = lib.get("records") or []
    if not records:
        sys.exit("library.json 里 records 是空的——没有文件可写")
    classified = lib.get("classified", True)

    # 分组。不分类时统一落到一个 sheet；分类时缺 category 的落到「未分类」而不是被悄悄丢掉
    groups, order = {}, []
    for r in records:
        cat = "全部文件" if not classified else (str(r.get("category") or "").strip() or UNCLASSIFIED)
        if cat not in groups:
            groups[cat] = []
            order.append(cat)
        groups[cat].append(r)
    if classified and lib.get("categories"):
        want = [c for c in lib["categories"] if c in groups]
        order = want + [c for c in order if c not in want]

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B6FD4")
    used = set()
    for cat in order:
        ws = wb.create_sheet(sheet_name(cat, used))
        ws.append(columns)
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for r in groups[cat]:
            ws.append([str(r.get(k, "") or "") for k in fields])
        for i, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 20)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"                 # 表头钉住：几十行往下翻还看得见列名
    wb.save(a.out)

    print(f"已生成 {a.out}：{len(order)} 个 sheet / {len(records)} 份")
    for cat in order:
        print(f"  · {cat}：{len(groups[cat])} 份")
    if not classified:
        print("  （本次未分类，全部放在一个 sheet）")
    print("提醒用户：在产出栏点开 library.xlsx 可以按 sheet 预览，"
          "并直接在预览里把某篇改到别的分类，或一键按分类归档到子文件夹。")


if __name__ == "__main__":
    main()
