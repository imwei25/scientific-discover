#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""扫描一个文件夹里的文件，抽出每份的开头内容，供后续由模型填台账与分类。

【这一步【不】判断年份/作者/杂志/核心观点/主要内容】——那是模型读 library_texts/*.txt 之后的事。
本脚本只做确定性的部分：列文件、抽文本、抽 PDF 自带的元数据、认 DOI，并把抽不动的如实标出来。

格式：默认只扫 `.pdf` / `.docx` / `.doc`（文献场景）。`--ext` 可以自选要整理的格式，
Markdown、纯文本以及用户自定义的任何**文本类**扩展名（.csv .json .srt .html …）都能读——
按【前若干行 / 若干字】截取，理由与 PDF 只读前几页一样：够填台账就行。
二进制格式（.xlsx .pptx .zip …）本脚本读不了，会如实标 err，别在台账里编内容。

用法：
  python scan_library.py                          # 扫当前目录（会话工作目录 = 用户选的那个文件夹）
  python scan_library.py --dir . --recursive       # 连子文件夹一起扫
  python scan_library.py --pages 4 --chars 6000    # 每份多抽一点（默认 3 页 / 4000 字）
  python scan_library.py --ext pdf,docx,md,txt     # 连 Markdown 与纯文本一起整理
  python scan_library.py --ext md,csv,srt --lines 80   # 只整理这几种文本文件，每份读前 80 行

产物（都写在 --dir 底下）：
  library_index.json   每份一条：文件名、大小、DOI、PDF 元数据、抽到多少字、正文文件名
  library_texts/*.txt  每份抽出来的开头内容（模型读这个填表，别去读原始文件）
"""
import argparse
import json
import os
import re
import sys
from datetime import datetime

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b")
# 本技能自己的产物 / 中间目录，二次扫描时必须跳过，否则会把上一轮的东西当成待整理文件
SKIP_DIRS = {"library_texts", ".preview", "__pycache__", ".git", "node_modules"}
DEFAULT_EXTS = "pdf,docx,doc"
WORD_EXTS = {".docx", ".doc"}
# 本技能自己产出的文件。开了 --ext md,txt 之后它们会混进扫描结果（上一轮的中间产物被当成
# "用户的文件"整理一遍），所以按文件名剔掉
SKIP_FILES = {"library_index.json", "library.json", "library.xlsx", "archive_log.json"}
# 明确读不了的二进制格式：与其按纯文本硬读出一堆乱码让模型去"理解"，不如如实报错
BINARY_EXTS = {".xlsx", ".xls", ".pptx", ".ppt", ".zip", ".rar", ".7z", ".png", ".jpg", ".jpeg",
               ".gif", ".bmp", ".tif", ".tiff", ".mp3", ".mp4", ".avi", ".mov", ".exe", ".dll"}


def _safe_stem(name, n):
    """给正文文件起名：序号 + 原名消毒（跨平台安全、不撞名）"""
    stem = os.path.splitext(os.path.basename(name))[0]
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", stem).strip() or "doc"
    return f"{n:03d}__{stem[:60]}"


def _pdf(path, pages, chars):
    """PDF：pymupdf 抽前若干页。顺带取 PDF 自带的 title/author（常常是空的或垃圾，只作参考）"""
    try:
        import fitz  # pymupdf
    except ImportError:
        return {"err": "缺少 pymupdf：先跑 env-setup 技能"}, ""
    try:
        doc = fitz.open(path)
    except Exception as e:
        return {"err": f"打不开这个 PDF：{type(e).__name__}: {e}"}, ""
    try:
        meta = doc.metadata or {}
        buf = []
        for i in range(min(pages, doc.page_count)):
            try:
                buf.append(doc.load_page(i).get_text("text"))
            except Exception as e:            # 单页坏了不该毁掉整篇
                buf.append(f"\n[第 {i + 1} 页抽取失败：{type(e).__name__}]\n")
        info = {"pdf_title": (meta.get("title") or "").strip(),
                "pdf_author": (meta.get("author") or "").strip(),
                "pages_total": doc.page_count}
        return info, "\n".join(buf)[:chars]
    finally:
        try:
            doc.close()
        except Exception:
            pass


def _docx(path, chars):
    """Word：python-docx 取段落 + 表格文字。.doc（老格式）python-docx 读不了，如实报错"""
    try:
        import docx  # python-docx
    except ImportError:
        return {"err": "缺少 python-docx：先跑 env-setup 技能"}, ""
    try:
        d = docx.Document(path)
    except Exception as e:
        return {"err": f"读不了这个 Word（.doc 老格式请先另存为 .docx）：{type(e).__name__}"}, ""
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for t in d.tables:                        # 摘要有时排在表格里
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
        if sum(len(x) for x in parts) > chars:
            break
    return {}, "\n".join(parts)[:chars]


def _plain(path, lines, chars):
    """Markdown / 纯文本 / 用户自定义的任何文本类格式：读前若干行。

    编码不是 UTF-8 的中文文本（记事本存的 GBK、老系统导出的东西）在国内是常态，
    直接 utf-8 解码会整份报错 —— 依次试几种编码，最后兜底 latin-1（一定能解，字可能花，
    但至少让模型看到结构而不是一句"读取失败"）。
    """
    try:
        raw = open(path, "rb").read(max(chars * 4, 65536))
    except Exception as e:
        return {"err": f"读不了这个文件：{type(e).__name__}: {e}"}, ""
    # 前 4KB 里有 \0 基本可以断定是二进制：与其吐乱码，不如如实说读不了
    if b"\0" in raw[:4096]:
        return {"err": "看起来是二进制格式（不是文本文件），本模块读不了它的内容"}, ""
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        text = raw.decode("latin-1", errors="replace")
    got = text.splitlines()[:lines]
    return {"lines_read": len(got)}, "\n".join(got)[:chars]


def _parse_exts(spec):
    """'pdf, .md ,TXT' → {'.pdf', '.md', '.txt'}（用户自定义格式那个框直接喂进来也认）"""
    out = set()
    for tok in re.split(r"[,;、\s]+", str(spec or "")):
        tok = tok.strip().lower().lstrip("*")
        if not tok:
            continue
        out.add(tok if tok.startswith(".") else "." + tok)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=".", help="文件所在文件夹（默认当前目录）")
    ap.add_argument("--recursive", action="store_true", help="连子文件夹一起扫（默认只扫这一层）")
    ap.add_argument("--ext", default=DEFAULT_EXTS,
                    help="要整理的格式，逗号分隔（默认 pdf,docx,doc）。md/txt 及任何文本类扩展名都能加")
    ap.add_argument("--pages", type=int, default=3, help="每份 PDF 抽前几页（默认 3）")
    ap.add_argument("--lines", type=int, default=200, help="每份文本文件读前几行（默认 200）")
    ap.add_argument("--chars", type=int, default=4000, help="每份最多抽多少字（默认 4000）")
    ap.add_argument("--max", type=int, default=800, help="最多处理多少份（防止误选到一个巨大的目录）")
    ap.add_argument("--out", default="library_index.json")
    ap.add_argument("--textdir", default="library_texts")
    a = ap.parse_args()

    root = os.path.abspath(a.dir)
    if not os.path.isdir(root):
        sys.exit(f"目录不存在：{root}")

    exts = _parse_exts(a.ext)
    if not exts:
        sys.exit("--ext 没解析出任何格式（例：--ext pdf,docx,md,txt）")
    binary_asked = sorted(e for e in exts if e in BINARY_EXTS)

    def want(fn):
        return os.path.splitext(fn)[1].lower() in exts and fn.lower() not in SKIP_FILES

    files = []
    if a.recursive:
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
            for fn in filenames:
                if want(fn):
                    files.append(os.path.join(dirpath, fn))
    else:
        for fn in os.listdir(root):
            p = os.path.join(root, fn)
            if os.path.isfile(p) and want(fn):
                files.append(p)
    files.sort(key=lambda p: os.path.relpath(p, root).lower())

    truncated = len(files) > a.max
    if truncated:
        files = files[:a.max]

    textdir = os.path.join(root, a.textdir)
    os.makedirs(textdir, exist_ok=True)

    records, n_ocr, n_err = [], 0, 0
    for i, p in enumerate(files, 1):
        rel = os.path.relpath(p, root).replace(os.sep, "/")
        ext = os.path.splitext(p)[1].lower()
        rec = {"id": i, "file": rel, "ext": ext.lstrip("."),
               "size": os.path.getsize(p) if os.path.exists(p) else 0}
        if ext == ".pdf":
            info, text = _pdf(p, a.pages, a.chars)
        elif ext in WORD_EXTS:
            info, text = _docx(p, a.chars)
        elif ext in BINARY_EXTS:
            # 用户在自定义格式框里填了 xlsx / pptx 这类：明说读不了，别按纯文本硬读出乱码
            info, text = {"err": f"{ext} 是二进制格式，本模块只能读文本类文件与 PDF/Word"}, ""
        else:
            # md / txt / 用户自定义的任何文本类扩展名走这条：读前若干行
            info, text = _plain(p, a.lines, a.chars)
        rec.update({k: v for k, v in info.items() if v not in ("", None)})
        text = (text or "").strip()
        rec["chars"] = len(text)
        if info.get("err"):
            n_err += 1
        elif len(text) < 200:
            # 抽出来几乎没字。PDF 里这几乎一定是【图片型扫描件】——不是"这篇是空的"，
            # 得让模型知道要走 OCR。Word 不可能是扫描件（它本来就是文字），所以只标"内容很少"，
            # 别把一篇短文档误报成扫描件让模型白跑一趟 OCR。
            if ext == ".pdf":
                rec["needs_ocr"] = True
                n_ocr += 1
            else:
                rec["sparse"] = True
        m = DOI_RE.search(text)
        if m:
            rec["doi"] = m.group(0).rstrip(".,;)")
        if text:
            tf = _safe_stem(rel, i) + ".txt"
            with open(os.path.join(textdir, tf), "w", encoding="utf-8") as f:
                f.write(f"# 文件：{rel}\n\n{text}\n")
            rec["text_file"] = f"{a.textdir}/{tf}"
        records.append(rec)

    out = {"scanned_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
           "dir": root, "recursive": bool(a.recursive),
           "exts": sorted(exts),
           "count": len(records), "truncated": truncated, "records": records}
    outp = os.path.join(root, a.out)
    with open(outp, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"扫描完成：{len(records)} 份（{root}）　格式：{'、'.join(sorted(exts))}")
    print(f"  内容抽取 → {a.textdir}/（每份一个 .txt，按它填表，别再去读原始文件）")
    print(f"  清单 → {a.out}")
    if n_ocr:
        print(f"  ⚠ {n_ocr} 份几乎抽不到文字（图片型扫描件）：needs_ocr=true，要么走 ocr 技能，"
              f"要么在表里如实写「原文未能识别」，不要凭文件名猜内容")
    if n_err:
        print(f"  ⚠ {n_err} 份读取失败：见 library_index.json 的 err 字段，如实报给用户")
    if binary_asked:
        print(f"  ⚠ 指定的格式里 {'、'.join(binary_asked)} 是二进制的，本模块读不了内容 —— "
              f"这些文件只会有文件名，如实告诉用户，别按文件名编「主要内容」")
    if truncated:
        print(f"  ⚠ 文件数超过 --max={a.max}，只处理了前 {a.max} 份——请告诉用户，别让他以为全做完了")
    if not records:
        print(f"  这个目录里没有 {'、'.join(sorted(exts))} 文件。请让用户确认选的文件夹、"
              f"以及要整理的格式（表单里可勾 Markdown / txt，或自己填扩展名）对不对。")


if __name__ == "__main__":
    main()
